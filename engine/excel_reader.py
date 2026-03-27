"""
PauseKreyol — Lecteur Excel
Lit les données clés des fichiers Association pour le dashboard et les alertes.
"""

from pathlib import Path
from openpyxl import load_workbook
from datetime import date, datetime
from typing import Optional
import json


def _find_asso_file(client_dir: Path) -> Optional[Path]:
    """Trouve le fichier ASSOCIATION_*.xlsx dans le dossier client."""
    matches = list(client_dir.glob("ASSOCIATION_*.xlsx"))
    return matches[0] if matches else None


def _parse_date(value) -> Optional[date]:
    """Convertit une valeur Excel en date Python."""
    if value is None:
        return None
    if isinstance(value, (date, datetime)):
        return value.date() if isinstance(value, datetime) else value
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return datetime.strptime(value, fmt).date()
            except ValueError:
                continue
    return None


def _alert_status(expiry_date: Optional[date]) -> str:
    """Calcule le statut d'alerte selon la date d'expiration."""
    if expiry_date is None:
        return "—"
    today = date.today()
    delta = (expiry_date - today).days
    if delta < 0:
        return "EXPIRÉ"
    elif delta < 90:
        return "URGENT"
    elif delta < 180:
        return "Attention"
    return "OK"


def read_client_summary(client_dir: Path) -> dict:
    """
    Lit les informations clés du fichier Association :
    - Identité (nom, SIRET, responsable)
    - Synthèse financière N-1 / N / N+1
    - Synthèse RH
    """
    asso_file = _find_asso_file(client_dir)
    if not asso_file:
        return {"error": "Fichier Association introuvable"}

    wb = load_workbook(asso_file, data_only=True)

    summary = {}

    # ── Identité ─────────────────────────────────────────────────
    if "IDENTITE" in wb.sheetnames:
        ws = wb["IDENTITE"]
        summary["identite"] = {
            "nom_officiel": ws["B5"].value,
            "nom_usuel": ws["B6"].value,
            "siret": ws["B7"].value,
            "president": ws["B16"].value,
            "email": ws["B21"].value,
            "telephone": ws["B22"].value,
            "adresse": ws["B12"].value,
            "licence_type1": ws["E6"].value,
            "licence_type2": ws["E7"].value,
            "licence_type3": ws["E8"].value,
        }

    # ── Finances N-1 / N / N+1 ───────────────────────────────────
    if "FINANCES" in wb.sheetnames:
        ws = wb["FINANCES"]
        summary["finances"] = {
            "total_produits": {
                "n1": ws["C19"].value,
                "n": ws["D19"].value,
                "n1_proj": ws["E19"].value,
            },
            "total_charges": {
                "n1": ws["C39"].value,
                "n": ws["D39"].value,
                "n1_proj": ws["E39"].value,
            },
            "resultat_net": {
                "n1": ws["C41"].value,
                "n": ws["D41"].value,
                "n1_proj": ws["E41"].value,
            },
            "tresorerie": {
                "n1": ws["C46"].value,
                "n": ws["D46"].value,
            },
        }

    # ── RH ───────────────────────────────────────────────────────
    if "RH_ASSO" in wb.sheetnames:
        ws = wb["RH_ASSO"]
        summary["rh"] = {
            "etpt_permanents": {
                "n1": ws["C14"].value,
                "n": ws["D14"].value,
                "n1_proj": ws["E14"].value,
            },
            "etpt_intermittents": {
                "n1": ws["C21"].value,
                "n": ws["D21"].value,
                "n1_proj": ws["E21"].value,
            },
        }

    # ── Collecte docs — progression ──────────────────────────────
    if "COLLECTE_DOCS" in wb.sheetnames:
        ws = wb["COLLECTE_DOCS"]
        total, recus = 0, 0
        for row in ws.iter_rows(min_row=5, max_row=40, min_col=4, max_col=4, values_only=True):
            val = row[0]
            if val is not None:
                total += 1
                if str(val).strip().lower() == "oui":
                    recus += 1
        summary["collecte_docs"] = {
            "total": total,
            "recus": recus,
            "pct": round(recus / total * 100) if total > 0 else 0,
        }

    # ── Projets dans le dossier ──────────────────────────────────
    meta_path = client_dir / "client.json"
    if meta_path.exists():
        meta = json.loads(meta_path.read_text())
        summary["projets"] = meta.get("projets", [])

    return summary


def read_alerts(client_dir: Path) -> list:
    """
    Lit l'onglet CAISSES du fichier Association et retourne
    la liste des alertes avec leur statut calculé.
    """
    asso_file = _find_asso_file(client_dir)
    if not asso_file:
        return []

    wb = load_workbook(asso_file, data_only=True)
    alerts = []

    # ── Alertes depuis CAISSES ───────────────────────────────────
    if "CAISSES" in wb.sheetnames:
        ws = wb["CAISSES"]
        # Lignes de données : col A = organisme, col C = N° adhérent, col D = date expiration
        for row in ws.iter_rows(min_row=5, max_row=49, values_only=True):
            organisme, ouvert, numero, date_val, a_jour, derniere_verif, _, contact = (
                row + (None,) * 8
            )[:8]

            if not organisme or organisme == organisme.upper() and len(str(organisme)) > 30:
                # C'est un titre de section, on skip
                if organisme and not any(c.isdigit() for c in str(organisme)):
                    continue

            if organisme and (numero or date_val):
                expiry = _parse_date(date_val)
                statut = _alert_status(expiry)
                alerts.append({
                    "organisme": str(organisme),
                    "numero": str(numero) if numero else None,
                    "date_expiration": expiry.isoformat() if expiry else None,
                    "statut": statut,
                    "a_jour": str(a_jour) if a_jour else None,
                    "contact": str(contact) if contact else None,
                })

    # ── Alertes licences depuis IDENTITE ─────────────────────────
    if "IDENTITE" in wb.sheetnames:
        ws = wb["IDENTITE"]
        licences = [
            ("Licence Type 1", ws["E6"].value, ws["E16"].value),
            ("Licence Type 2", ws["E7"].value, ws["E17"].value),
            ("Licence Type 3", ws["E8"].value, ws["E18"].value),
            ("Agrément JEP", ws["E11"].value, ws["E19"].value),
            ("QUALIOPI", ws["E12"].value, ws["E20"].value),
            ("RC Pro", None, ws["E21"].value),
            ("Mutuelle obligatoire", None, ws["E22"].value),
        ]
        for label, numero, date_val in licences:
            if numero or date_val:
                expiry = _parse_date(date_val)
                statut = _alert_status(expiry)
                alerts.append({
                    "organisme": label,
                    "numero": str(numero) if numero else None,
                    "date_expiration": expiry.isoformat() if expiry else None,
                    "statut": statut,
                    "a_jour": None,
                    "contact": None,
                })

    return alerts
