"""
PauseKreyol — Parseur de fiche client V2
Lit la FICHE_CLIENT_V2 et extrait toutes les données structurées.
Détecte automatiquement V1 et V2 du template.
"""

import logging
from pathlib import Path
from typing import Optional
from openpyxl import load_workbook

logger = logging.getLogger(__name__)


def _val(ws, ref) -> Optional[str]:
    """Lit une cellule et retourne None si vide ou placeholder."""
    v = ws[ref].value
    if v is None:
        return None
    s = str(v).strip()
    # Ignore les placeholders en italique gris
    if s in ("", "—", "-", "JJ/MM/AAAA") or s.startswith("Cocher les cases") \
       or s.startswith("Séparez par") or s.startswith("Droits de") \
       or s.startswith("Statuts") or "Obligatoire pour" in s:
        return None
    return s


def detect_template_type(wb) -> str:
    """Détecte le type de template."""
    sheets = wb.sheetnames
    if "FICHE_CLIENT" in sheets:
        ws = wb["FICHE_CLIENT"]
        marker = str(ws["A3"].value or "")
        if "V2" in marker:
            return "intake_v2"
        if "PAUSEKREYOL_INTAKE_FORM" in marker:
            return "intake_v1"
    if "ASSO_FRANCE" in sheets:
        return "collecte_france"
    if "SAS_SENEGAL" in sheets:
        return "collecte_senegal"
    if "IDENTITE" in sheets and "CAISSES" in sheets:
        return "structure_fr"
    if "IDENTITE" in sheets and "ORGANISMES" in sheets:
        return "structure_sn"
    return "unknown"


def _determine_type_structure(statut: str, implantation: str) -> str:
    """Déduit le type de structure depuis statut + implantation."""
    if not statut and not implantation:
        return "asso_france"
    statut_l = (statut or "").lower()
    impl_l = (implantation or "").lower()
    if "sénégal" in impl_l or "senegal" in impl_l or "afrique" in impl_l:
        return "sas_senegal"
    if "antilles" in impl_l or "martinique" in impl_l or "guadeloupe" in impl_l:
        return "asso_antilles"
    if "sas" in statut_l or "sarl" in statut_l or "sa " in statut_l:
        return "sas_france"
    return "asso_france"


def parse_intake_v2(wb) -> dict:
    """Parse FICHE_CLIENT V2 (nouveau format avec statut/implantation séparés)."""
    ws = wb["FICHE_CLIENT"]

    # ── Section 1 — Type de structure ──────────────────────────────────────
    statut_juridique  = _val(ws, "B5")
    implantation      = _val(ws, "E5")
    categorie         = _val(ws, "B6")   # Créateur / Diffuseur / Créateur & Diffuseur
    marche_cible      = _val(ws, "E6")
    domaines          = _val(ws, "B7")   # texte libre avec cases cochées
    innovation        = _val(ws, "B8")
    nb_projets_an     = _val(ws, "E8")

    type_structure = _determine_type_structure(statut_juridique, implantation)

    # ── Section 2 — Identité ────────────────────────────────────────────────
    nom_officiel      = _val(ws, "B10")
    nom_usuel         = _val(ws, "E10")
    siret_ninea       = _val(ws, "B11")
    code_ape          = _val(ws, "E11")
    date_creation     = _val(ws, "B12")
    rna_rccm          = _val(ws, "E12")
    date_ag           = _val(ws, "B13")
    clotures_ok       = _val(ws, "E13")
    adresse_siege     = _val(ws, "B14")
    ville             = _val(ws, "E14")
    code_postal       = _val(ws, "B15")
    pays              = _val(ws, "E15") or "France"

    # ── Section 3 — Contacts ────────────────────────────────────────────────
    president         = _val(ws, "B17")
    email_contact     = _val(ws, "E17")
    tresorier         = _val(ws, "B18")
    telephone         = _val(ws, "E18")
    secretaire        = _val(ws, "B19")
    site_internet     = _val(ws, "E19")
    dir_artistique    = _val(ws, "B20")
    reseaux           = _val(ws, "E20")
    emails_agent      = _val(ws, "B21")  # emails surveillance agent IA

    # ── Section 4 — Licences ────────────────────────────────────────────────
    licence_statut    = _val(ws, "B23")
    licence_numero    = _val(ws, "E23")
    date_expir_lic    = _val(ws, "B24")
    numero_urssaf     = _val(ws, "E24")
    numero_guso       = _val(ws, "B25")
    numero_audiens    = _val(ws, "E25")
    numero_afdas      = _val(ws, "B26")
    numero_conges     = _val(ws, "E26")
    mutuelle          = _val(ws, "B27")
    numero_aem        = _val(ws, "E27")

    # ── Section 5 — Droits d'auteur ─────────────────────────────────────────
    societe_droits    = _val(ws, "B29")
    num_societaire    = _val(ws, "E29")
    types_droits      = _val(ws, "B30")

    # ── Section 6 — Banque ──────────────────────────────────────────────────
    banque            = _val(ws, "B32")
    iban              = _val(ws, "E32")
    bic               = _val(ws, "B33")
    titulaire_compte  = _val(ws, "E33")

    # ── Section 7 — Documents ───────────────────────────────────────────────
    docs = {
        "statuts":          _val(ws, "B35"),
        "pv_ag":            _val(ws, "E35"),
        "bilan_n1":         _val(ws, "B36"),
        "compte_result":    _val(ws, "E36"),
        "rib":              _val(ws, "B37"),
        "licence_pdf":      _val(ws, "E37"),
        "attestation_urs":  _val(ws, "B38"),
        "dossier_artis":    _val(ws, "E38"),
    }

    # ── Construction client_data ────────────────────────────────────────────
    # Détermine siret/ninea selon implantation
    is_senegal = "sas_senegal" in type_structure
    adresse_complete = " ".join(filter(None, [adresse_siege, code_postal, ville]))

    client_data = {
        "type_structure":       type_structure,
        "statut_juridique":     statut_juridique,
        "implantation":         implantation,
        "categorie":            categorie,
        "marche_cible":         marche_cible,
        "domaines_artistiques": domaines,
        "innovation":           innovation,

        "nom_officiel":         nom_officiel,
        "nom_usuel":            nom_usuel,
        "siret":                siret_ninea if not is_senegal else None,
        "ninea":                siret_ninea if is_senegal else None,
        "code_ape":             code_ape,
        "date_creation":        date_creation,
        "numero_rna":           rna_rccm if not is_senegal else None,
        "rccm":                 rna_rccm if is_senegal else None,
        "date_derniere_ag":     date_ag,
        "clotures_comptables":  clotures_ok,
        "adresse_siege":        adresse_complete or adresse_siege,
        "pays":                 pays,

        "president":            president,
        "tresorier":            tresorier,
        "secretaire":           secretaire,
        "directeur_artistique": dir_artistique,
        "email_contact":        email_contact,
        "telephone":            telephone,
        "site_internet":        site_internet,
        "reseaux_sociaux":      reseaux,
        "emails_surveillance":  [e.strip() for e in emails_agent.split(",") if e.strip()] if emails_agent else [],

        "licence_statut":       licence_statut,
        "licence_type1":        licence_numero,
        "date_expiration_licence": date_expir_lic,
        "numero_urssaf":        numero_urssaf,
        "numero_guso":          numero_guso,
        "numero_audiens":       numero_audiens,
        "numero_afdas":         numero_afdas,
        "numero_conges_spectacles": numero_conges,
        "mutuelle":             mutuelle,
        "numero_aem":           numero_aem,

        "societe_droits":       societe_droits,
        "num_societaire":       num_societaire,
        "types_droits":         types_droits,

        "banque":               banque,
        "iban":                 iban,
        "bic":                  bic,
        "titulaire_compte":     titulaire_compte,

        "documents_fournis":    {k: v for k, v in docs.items() if v},
    }

    # Nettoyage : supprime les None
    client_data = {k: v for k, v in client_data.items() if v is not None}

    return client_data


def parse_intake_v1(wb) -> dict:
    """Parse ancien format V1 pour rétrocompatibilité."""
    ws = wb["FICHE_CLIENT"]
    type_raw = (ws["D6"].value or "").lower()
    type_structure = "sas_senegal" if any(x in type_raw for x in ["sas","sénégal","senegal","ohada"]) else "asso_france"
    identifiant = str(ws["D7"].value or "").strip()
    rna_rccm = str(ws["D8"].value or "").strip()
    return {
        "type_structure": type_structure,
        "nom_officiel":   str(ws["B6"].value or "").strip() or None,
        "nom_usuel":      str(ws["B7"].value or "").strip() or None,
        "date_creation":  str(ws["B8"].value or "").strip() or None,
        "adresse_siege":  str(ws["B9"].value or "").strip() or None,
        "pays":           str(ws["D10"].value or "France").strip(),
        "siret":          identifiant if type_structure == "asso_france" else None,
        "ninea":          identifiant if type_structure == "sas_senegal" else None,
        "numero_rna":     rna_rccm if type_structure == "asso_france" else None,
        "rccm":           rna_rccm if type_structure == "sas_senegal" else None,
        "president":      str(ws["B13"].value or "").strip() or None,
        "email_contact":  str(ws["D13"].value or "").strip() or None,
        "telephone":      str(ws["D14"].value or "").strip() or None,
        "licence_type1":  str(ws["B28"].value or "").strip() or None,
    }


def parse_excel_client(file_path: Path) -> dict:
    """
    Point d'entrée principal. Parse n'importe quel Excel client.
    Retourne un dict standardisé avec client_data + projet_data + meta.
    """
    result = {
        "type_fichier": "unknown",
        "type_structure": "asso_france",
        "client_data": {},
        "projet_data": {},
        "nb_champs": 0,
        "champs_manquants": [],
        "erreur": None,
    }

    try:
        wb = load_workbook(str(file_path), data_only=True)
        ttype = detect_template_type(wb)
        result["type_fichier"] = ttype

        if ttype == "intake_v2":
            client_data = parse_intake_v2(wb)
        elif ttype == "intake_v1":
            client_data = parse_intake_v1(wb)
        elif ttype == "collecte_france":
            client_data = _parse_collecte(wb, "ASSO_FRANCE", "asso_france")
        elif ttype == "collecte_senegal":
            client_data = _parse_collecte(wb, "SAS_SENEGAL", "sas_senegal")
        elif ttype in ("structure_fr", "structure_sn"):
            client_data = _parse_structure(wb, ttype)
        else:
            result["erreur"] = "Format non reconnu. Utilisez TEMPLATE_FICHE_CLIENT_V2.xlsx"
            return result

        result["type_structure"] = client_data.get("type_structure", "asso_france")
        result["client_data"] = client_data
        result["nb_champs"] = len([v for v in client_data.values() if v])

        # Champs obligatoires
        required = ["nom_officiel", "email_contact", "president"]
        is_sn = result["type_structure"] == "sas_senegal"
        required.append("ninea" if is_sn else "siret")
        result["champs_manquants"] = [f for f in required if not client_data.get(f)]

        # Licence obligatoire si diffuseur
        cat = (client_data.get("categorie") or "").lower()
        if "diffuseur" in cat and not client_data.get("licence_type1"):
            result["champs_manquants"].append("licence_type1 (obligatoire pour diffuseur)")

        # Droits d'auteur si écrivain
        domaines = (client_data.get("domaines_artistiques") or "").lower()
        if "écriture" in domaines or "écrivain" in domaines or "littérature" in domaines:
            if not client_data.get("societe_droits"):
                result["champs_manquants"].append("societe_droits (recommandé pour auteur)")

        logger.info(f"Parse {ttype} : {result['nb_champs']} champs, {len(result['champs_manquants'])} manquants")

    except Exception as e:
        result["erreur"] = str(e)
        logger.error(f"Erreur parsing {file_path} : {e}")

    return result


def _parse_collecte(wb, sheet, type_structure) -> dict:
    """Parse les anciens templates COLLECTE_CLIENT."""
    ws = wb[sheet]
    if sheet == "ASSO_FRANCE":
        return {
            "type_structure": type_structure,
            "nom_officiel": _val(ws, "B5"), "nom_usuel": _val(ws, "B6"),
            "siret": _val(ws, "B7"), "code_ape": _val(ws, "B8"),
            "numero_rna": _val(ws, "B9"), "date_creation": _val(ws, "B10"),
            "adresse_siege": _val(ws, "B11"), "president": _val(ws, "B14"),
            "tresorier": _val(ws, "B15"), "email_contact": _val(ws, "B18"),
            "telephone": _val(ws, "B19"), "licence_type1": _val(ws, "B23"),
        }
    else:
        return {
            "type_structure": type_structure,
            "nom_officiel": _val(ws, "B5"), "ninea": _val(ws, "B7"),
            "rccm": _val(ws, "B8"), "president": _val(ws, "B15"),
            "email_contact": _val(ws, "B17"), "telephone": _val(ws, "B18"),
        }


def _parse_structure(wb, ttype) -> dict:
    """Parse les fichiers STRUCTURE existants."""
    ws = wb["IDENTITE"]
    is_sn = ttype == "structure_sn"
    return {
        "type_structure": "sas_senegal" if is_sn else "asso_france",
        "nom_officiel": _val(ws, "B5") or _val(ws, "B4"),
        "nom_usuel": _val(ws, "B6") or _val(ws, "B5"),
        "siret": _val(ws, "B7") if not is_sn else None,
        "ninea": _val(ws, "B6") if is_sn else None,
        "president": _val(ws, "B16") if not is_sn else _val(ws, "B15"),
        "email_contact": _val(ws, "B21") if not is_sn else _val(ws, "B17"),
        "telephone": _val(ws, "B22") if not is_sn else _val(ws, "B18"),
        "adresse_siege": _val(ws, "B12") if not is_sn else _val(ws, "B10"),
        "licence_type1": _val(ws, "E6") if not is_sn else None,
    }


def diff_with_existing(parsed: dict, existing_client_data: dict) -> dict:
    """Compare données parsées avec dossier existant. Retourne les écarts."""
    changes = {}
    new_data = parsed.get("client_data", {})
    for key, new_val in new_data.items():
        if key in ("type_structure", "documents_fournis", "emails_surveillance"):
            continue
        old_val = existing_client_data.get(key)
        if new_val and str(new_val).strip() != str(old_val or "").strip():
            changes[key] = {"avant": old_val, "après": new_val}
    return changes
