"""
PauseKreyol — Moteur de templates
Crée automatiquement un dossier client avec les Excel pré-remplis et interconnectés.
"""

import shutil
import json
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook

BASE_DIR = Path(__file__).parent.parent
TEMPLATES_DIR = BASE_DIR / "templates"
CLIENTS_DIR = BASE_DIR / "clients"

# ─────────────────────────────────────────────
# Mapping des cellules à pré-remplir
# Format : { "NomSheet": { "CellRef": "clé_dans_client_data" } }
# ─────────────────────────────────────────────

ASSOCIATION_MAPPING = {
    "IDENTITE": {
        # Col B = valeurs (label en col A)
        "B5": "nom_officiel",
        "B6": "nom_usuel",
        "B7": "siret",
        "B8": "code_ape",
        "B9": "numero_rna",
        "B10": "date_creation",
        "B12": "adresse_siege",
        # Direction (label col A, valeur col B)
        "B16": "president",
        "B17": "tresorier",
        "B19": "directeur_artistique",
        "B21": "email_contact",
        "B22": "telephone",
        # Licences (label col D, valeur col E)
        "E6": "licence_type1",
        "E7": "licence_type2",
        "E8": "licence_type3",
    }
}

BUDGET_MAPPING = {
    "PARAMETRES": {
        "B4": "nom_projet",
        "B5": "nom_officiel",
        "B6": "siret",
        "B7": "licence_type1",
        "B8": "date_debut_projet",
        "B9": "date_fin_projet",
        "B10": "lieu_projet",
        "B11": "code_aap",
        "B14": "president",
        "B16": "email_contact",
        "B17": "telephone",
        "B18": "adresse_siege",
    }
}


def _fill_workbook(wb, mapping: dict, data: dict) -> int:
    """Remplit les cellules d'un workbook selon le mapping. Retourne le nb de cellules remplies."""
    filled = 0
    for sheet_name, cells in mapping.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        for cell_ref, data_key in cells.items():
            value = data.get(data_key)
            if value is not None and value != "":
                ws[cell_ref] = value
                filled += 1
    return filled


def create_client_folder(client_data: dict) -> Path:
    """
    Crée le dossier complet d'un nouveau client.

    client_data doit contenir au minimum :
        - nom_officiel : str
        - siret : str

    Retourne le chemin du dossier créé.
    """
    nom = client_data.get("nom_usuel") or client_data.get("nom_officiel", "NOUVEAU_CLIENT")
    # Sanitize nom pour le filesystem
    slug = "".join(c if c.isalnum() or c in " _-" else "_" for c in nom).strip().replace(" ", "_")
    timestamp = datetime.now().strftime("%Y%m")

    client_dir = CLIENTS_DIR / f"{slug}_{timestamp}"
    projets_dir = client_dir / "projets"
    docs_dir = client_dir / "documents"

    client_dir.mkdir(parents=True, exist_ok=True)
    projets_dir.mkdir(exist_ok=True)
    docs_dir.mkdir(exist_ok=True)

    # ── 1. Fichier Association ──────────────────────────────────────
    asso_src = TEMPLATES_DIR / "TEMPLATE_ASSOCIATION.xlsx"
    asso_dst = client_dir / f"ASSOCIATION_{slug}.xlsx"
    shutil.copy2(asso_src, asso_dst)

    wb_asso = load_workbook(asso_dst)
    filled = _fill_workbook(wb_asso, ASSOCIATION_MAPPING, client_data)
    wb_asso.save(asso_dst)
    print(f"  ✅ Association : {asso_dst.name} ({filled} cellules pré-remplies)")

    # ── 2. Fichier Budget projet vierge ─────────────────────────────
    budget_src = TEMPLATES_DIR / "TEMPLATE_BUDGET_PROJET.xlsx"
    budget_dst = projets_dir / f"BUDGET_PROJET_VIERGE_{slug}.xlsx"
    shutil.copy2(budget_src, budget_dst)

    wb_budget = load_workbook(budget_dst)
    filled = _fill_workbook(wb_budget, BUDGET_MAPPING, client_data)
    wb_budget.save(budget_dst)
    print(f"  ✅ Budget projet : {budget_dst.name} ({filled} cellules pré-remplies)")

    # ── 3. Fichier metadata client (JSON) ───────────────────────────
    meta = {
        "slug": slug,
        "created_at": datetime.now().isoformat(),
        "client_data": client_data,
        "fichiers": {
            "association": str(asso_dst.relative_to(BASE_DIR)),
            "budget_template": str(budget_dst.relative_to(BASE_DIR)),
        },
        "projets": []
    }
    meta_path = client_dir / "client.json"
    meta_path.write_text(json.dumps(meta, ensure_ascii=False, indent=2))
    print(f"  ✅ Metadata : {meta_path.name}")

    # ── 4. Sync vers Google Drive (prod uniquement) ─────────────────
    try:
        from engine.drive_storage import sync_client_to_drive
        meta = sync_client_to_drive(client_dir, meta)
        meta_path.write_text(json.dumps(meta, ensure_ascii=False, indent=2))
        print(f"  ✅ Synchronisé sur Google Drive")
    except Exception as e:
        print(f"  ⚠️  Drive sync ignorée : {e}")

    return client_dir


def create_project(client_slug: str, project_data: dict) -> Path:
    """
    Crée un nouveau dossier projet pour un client existant.

    project_data doit contenir :
        - nom_projet : str
        - date_debut_projet : str (optionnel)
        - date_fin_projet : str (optionnel)
        - lieu_projet : str (optionnel)
        - code_aap : str (optionnel)
    """
    # Trouver le dossier client
    matches = list(CLIENTS_DIR.glob(f"{client_slug}*"))
    if not matches:
        raise ValueError(f"Client '{client_slug}' non trouvé dans {CLIENTS_DIR}")
    client_dir = matches[0]

    # Charger les données client existantes
    meta_path = client_dir / "client.json"
    meta = json.loads(meta_path.read_text())
    client_data = meta["client_data"]

    # Merge : données client + données projet
    merged = {**client_data, **project_data}

    nom_projet = project_data.get("nom_projet", "PROJET")
    slug_projet = "".join(c if c.isalnum() or c in " _-" else "_" for c in nom_projet).replace(" ", "_")

    projet_dir = client_dir / "projets" / slug_projet
    projet_dir.mkdir(parents=True, exist_ok=True)

    # Copier et pré-remplir le budget
    budget_src = TEMPLATES_DIR / "TEMPLATE_BUDGET_PROJET.xlsx"
    budget_dst = projet_dir / f"BUDGET_{slug_projet}.xlsx"
    shutil.copy2(budget_src, budget_dst)

    wb = load_workbook(budget_dst)
    filled = _fill_workbook(wb, BUDGET_MAPPING, merged)
    wb.save(budget_dst)
    print(f"  ✅ Projet '{nom_projet}' : {budget_dst.name} ({filled} cellules pré-remplies)")

    # Mettre à jour le metadata client
    meta["projets"].append({
        "nom": nom_projet,
        "slug": slug_projet,
        "created_at": datetime.now().isoformat(),
        "fichier_budget": str(budget_dst.relative_to(BASE_DIR)),
        "statut": "en_construction"
    })
    meta_path.write_text(json.dumps(meta, ensure_ascii=False, indent=2))

    return projet_dir


def list_clients() -> list:
    """Retourne la liste de tous les clients."""
    clients = []
    for client_dir in CLIENTS_DIR.iterdir():
        meta_path = client_dir / "client.json"
        if meta_path.exists():
            meta = json.loads(meta_path.read_text())
            clients.append({
                "slug": meta["slug"],
                "nom": meta["client_data"].get("nom_usuel") or meta["client_data"].get("nom_officiel"),
                "siret": meta["client_data"].get("siret"),
                "nb_projets": len(meta.get("projets", [])),
                "created_at": meta["created_at"],
                "dossier": str(client_dir),
            })
    return clients


# ─────────────────────────────────────────────
# CLI rapide pour tester
# ─────────────────────────────────────────────
if __name__ == "__main__":
    import sys

    print("\n🎭 PauseKreyol — Moteur de templates\n")

    # Exemple : créer House of Bastet
    hob_data = {
        "nom_officiel": "House of Bastet",
        "nom_usuel": "House of Bastet (HOB)",
        "siret": "93280087300014",
        "code_ape": "9001Z",
        "president": "Madame Djenabou CISSE",
        "email_contact": "cie.underground.dance.providers@gmail.com",
        "telephone": "+33 06 58 91 08 14",
        "adresse_siege": "20 rue Primo Levy, 93000 Bobigny",
        "licence_type1": "PLATESV-D-2025-001444",
        "date_creation": "2024-09-01",
    }

    print("📁 Création du dossier client : House of Bastet")
    client_dir = create_client_folder(hob_data)
    print(f"\n📂 Dossier créé : {client_dir}\n")

    print("📋 Création d'un projet : Bastet Festival #5")
    projet_dir = create_project("House_of_Bastet", {
        "nom_projet": "Bastet Festival #5 — Oct 2026",
        "date_debut_projet": "2026-10-22",
        "date_fin_projet": "2026-10-24",
        "lieu_projet": "La Place, Paris",
        "code_aap": "HHP26S2",
    })
    print(f"\n📂 Projet créé : {projet_dir}\n")

    print("📊 Clients existants :")
    for c in list_clients():
        print(f"  • {c['nom']} — SIRET {c['siret']} — {c['nb_projets']} projet(s)")
